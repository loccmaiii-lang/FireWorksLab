# -*- coding: utf-8 -*-
"""分享 PPT 的设计令牌与组件（1920×1080 设计像素 → 16:9 PPTX）。

模板：协作/提示词/分享PPT交接_2026-10-29/模板_VFX组PPT风格.md；数值按交接第 3 节落到 1920×1080。
1 设计像素 = 6350 EMU；字号 px × 0.5 = pt。
"""
import re

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- 设计令牌 ----------
BG = '0A0A0A'        # 背景
ACC = '0FF796'       # 强调（小标签、序号、关键数字、细条），面积 ≤ 10%
CARD = '161616'      # 卡片底
TXT = 'F5F5F5'       # 文字
A_BODY, A_SUB, A_NOTE, A_FOOT = 1.0, 0.70, 0.55, 0.40   # 正文 / 次要 / 说明 / 页脚
A_STROKE = 0.08      # 卡片描边（白）
A_GREEN_BG = 0.08    # 弱绿底
A_GREEN_LINE = 0.35  # 绿描边
FONT_CN = '微软雅黑'
FONT_EN = 'Consolas'

S_TITLE, S_CARD, S_BODY, S_SMALL, S_LABEL, S_FOOT = 52, 30, 28, 24, 24, 24
L, R, W_CONTENT = 96, 1824, 1728     # 左右页边距 96
COL60_W, COL40_X, COL40_W = 1008, 1152, 672
Y_LABEL, Y_TITLE, Y_CONTENT, Y_CONTENT_END = 60, 102, 210, 980

EMU_PX = 6350


def E(v):
    return Emu(int(round(v * EMU_PX)))


def _rgb(h):
    return RGBColor.from_string(h)


def _add_alpha(srgb_el, alpha):
    if alpha is None or alpha >= 0.999:
        return
    for old in srgb_el.findall(qn('a:alpha')):
        srgb_el.remove(old)
    a = etree.SubElement(srgb_el, qn('a:alpha'))
    a.set('val', str(int(round(alpha * 100000))))


def _strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


class Deck:
    def __init__(self, footer):
        self.prs = Presentation()
        self.prs.slide_width = E(1920)
        self.prs.slide_height = E(1080)
        self.layout = self.prs.slide_layouts[6]
        self.footer = footer

    def slide(self, label=None, title=None, footer=True, transition=True):
        s = self.prs.slides.add_slide(self.layout)
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = _rgb(BG)
        if label:
            page_label(s, label)
        if title:
            page_title(s, title)
        if footer:
            page_footer(s, self.footer)
        if transition:
            add_fade(s)
        return s

    def save(self, path, title='', subject=''):
        cp = self.prs.core_properties
        cp.title, cp.subject = title, subject
        cp.author = cp.last_modified_by = ''
        self.prs.save(path)


def add_fade(slide):
    sld = slide._element
    for old in sld.findall(qn('p:transition')):
        sld.remove(old)
    tr = etree.Element(qn('p:transition'))
    tr.set('spd', 'med')
    etree.SubElement(tr, qn('p:fade'))
    clr = sld.find(qn('p:clrMapOvr'))
    if clr is not None:
        clr.addnext(tr)
    else:
        sld.append(tr)


# ---------- 基本形状 ----------
def rect(slide, x, y, w, h, fill=None, fill_alpha=None, line=None, line_alpha=None, line_w=1,
         dash=False, shape=MSO_SHAPE.RECTANGLE, name=None):
    sp = slide.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    _strip_style(sp)
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
        _add_alpha(sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr')), fill_alpha)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = E(line_w)
        ln = sp._element.spPr.find(qn('a:ln'))
        _add_alpha(ln.find(qn('a:solidFill')).find(qn('a:srgbClr')), line_alpha)
        if dash:
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sp.line.fill.background()
    if name:
        sp.name = name
    return sp


def line(slide, x1, y1, x2, y2, color=ACC, alpha=None, w=2, dash=False, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _strip_style(c)
    c.line.color.rgb = _rgb(color)
    c.line.width = E(w)
    ln = c._element.spPr.find(qn('a:ln'))
    _add_alpha(ln.find(qn('a:solidFill')).find(qn('a:srgbClr')), alpha)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle')
        te.set('w', 'med')
        te.set('len', 'med')
    return c


# ---------- 文字 ----------
def _style_run(run, size, color, alpha, bold, spc, font_en, font_cn):
    f = run.font
    f.size = Pt(size * 0.5)
    f.bold = bool(bold)
    f.color.rgb = _rgb(color)
    f.name = font_en
    rPr = run._r.get_or_add_rPr()
    rPr.set('lang', 'zh-CN')
    rPr.set('altLang', 'en-US')
    if spc:
        rPr.set('spc', str(int(round(spc * 0.5 * 100))))
    latin = rPr.find(qn('a:latin'))
    for tag in ('a:ea', 'a:cs'):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    ea = etree.Element(qn('a:ea'))
    ea.set('typeface', font_cn)
    cs = etree.Element(qn('a:cs'))
    cs.set('typeface', font_cn)
    latin.addnext(ea)
    ea.addnext(cs)
    _add_alpha(rPr.find(qn('a:solidFill')).find(qn('a:srgbClr')), alpha)


def text(slide, x, y, w, h, paras, size=S_BODY, color=TXT, alpha=A_BODY, bold=False, align='l',
         anchor='t', lh=1.6, spc=0, font_en=FONT_EN, font_cn=FONT_CN, space_after=0, name=None):
    """paras: str | [para]；para: str | [run]；run: str | dict(t=, color=, alpha=, bold=, size=, en=, cn=, spc=)。
    行高用固定磅值（字号 × lh），对应模板里的 CSS 行高。"""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        runs = [para] if isinstance(para, (str, dict)) else para
        psize = max((r.get('size', size) if isinstance(r, dict) else size) for r in runs) if runs else size
        p.line_spacing = Pt(psize * 0.5 * lh)
        if space_after and i < len(paras) - 1:
            p.space_after = Pt(space_after * 0.5)
        for rd in runs:
            if isinstance(rd, str):
                rd = {'t': rd}
            # 空格单独成段并用中文字体：Consolas 的空格有半个字宽，夹在中文和数字之间显得松
            # \v 表示段内换行（手动断行，避免词被拆到两行）
            for k, seg in enumerate(rd['t'].split('\v')):
                if k:
                    p.add_line_break()
                for chunk in re.split(r'( +)', seg):
                    if not chunk:
                        continue
                    r = p.add_run()
                    r.text = chunk
                    en = font_cn if chunk.strip() == '' else rd.get('en', font_en)
                    _style_run(r, rd.get('size', size), rd.get('color', color), rd.get('alpha', alpha),
                               rd.get('bold', bold), rd.get('spc', spc), en, rd.get('cn', font_cn))
    if name:
        tb.name = name
    return tb


# ---------- 页面骨架 ----------
def page_label(slide, label):
    text(slide, L, Y_LABEL, 1200, 36, label, size=S_LABEL, color=ACC, bold=True, spc=6, lh=1.3,
         name='label')


def page_title(slide, title):
    text(slide, L, Y_TITLE, W_CONTENT, 72, title, size=S_TITLE, color=TXT, bold=True, lh=1.25, name='title')


def page_footer(slide, footer):
    rect(slide, L, 1026, 18, 18, fill=ACC, name='footer_mark')
    text(slide, L + 34, 1015, 1200, 40, footer, size=S_FOOT, alpha=A_FOOT, lh=1.6, anchor='m', name='footer')


# ---------- 组件 ----------
def card(slide, x, y, w, h, key=False, fill=CARD):
    c = rect(slide, x, y, w, h, fill=fill, line='F5F5F5', line_alpha=A_STROKE, line_w=1)
    if key:
        rect(slide, x, y, 6, h, fill=ACC)
    return c


def end_bar(slide, x, y, w=54):
    return rect(slide, x, y, w, 6, fill=ACC)


def pill(slide, x, y, w, h, label, size=S_SMALL):
    rect(slide, x, y, w, h, fill=ACC, fill_alpha=0.10, line=ACC, line_alpha=A_GREEN_LINE, line_w=1)
    text(slide, x, y, w, h, label, size=size, color=ACC, bold=True, align='c', anchor='m', lh=1.2)


def placeholder(slide, x, y, w, h, head, sub, bg_image=None, bg_alpha=0.78):
    if bg_image:
        slide.shapes.add_picture(bg_image, E(x), E(y), E(w), E(h))
        rect(slide, x, y, w, h, fill=BG, fill_alpha=bg_alpha)
    else:
        rect(slide, x, y, w, h, fill=CARD)
    rect(slide, x, y, w, h, line=ACC, line_alpha=0.6, line_w=3, dash=True)
    text(slide, x + 40, y + h / 2 - 70, w - 80, 60, head, size=36, color=ACC, bold=True, align='c', anchor='m',
         lh=1.3)
    text(slide, x + 60, y + h / 2 + 4, w - 120, 90, sub, size=S_SMALL, alpha=A_NOTE, align='c', lh=1.6)


def image(slide, path, x, y, w, h, tag=None, tag_pos='bl'):
    pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    if tag:
        source_tag(slide, x, y, w, h, tag, tag_pos)
    return pic


def source_tag(slide, x, y, w, h, tag, pos='bl'):
    # 估算宽度：中文字宽约 1 em，Consolas 约 0.55 em，空格按中文字体约 0.3 em
    tw = 24 * sum(1.0 if ord(ch) > 0x2E80 else (0.3 if ch == ' ' else 0.55) for ch in tag) + 32
    th = 40
    tx = x + 12 if pos.endswith('l') else x + w - tw - 12
    ty = y + h - th - 12 if pos.startswith('b') else y + 12
    rect(slide, tx, ty, tw, th, fill=BG, fill_alpha=0.72)
    text(slide, tx, ty, tw, th, tag, size=S_SMALL, alpha=A_SUB, align='c', anchor='m', lh=1.2, name='source_tag')


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# ---------- 第 2 版新增组件 ----------
def before_after(slide, x, y, w, h, img_before, img_after, tag_before='修前', tag_after='修后', gap=12, src=None):
    """修前 / 修后成对图：等宽两格，左上角小标签，来源标注放右下角。"""
    cw = (w - gap) / 2
    for i, (img, tg) in enumerate(((img_before, tag_before), (img_after, tag_after))):
        cx = x + i * (cw + gap)
        slide.shapes.add_picture(img, E(cx), E(y), E(cw), E(h))
        tw = 24 * len(tg) + 28
        rect(slide, cx, y, tw, 38, fill=BG, fill_alpha=0.82)
        text(slide, cx, y, tw, 38, tg, size=S_SMALL, color=TXT if i == 0 else ACC, bold=True, align='c', anchor='m',
             lh=1.2)
    if src:
        source_tag(slide, x, y, w, h, src, pos='br')


def node_circle(slide, cx, cy, d, label, size=36):
    rect(slide, cx - d / 2, cy - d / 2, d, d, fill=ACC, fill_alpha=0.06, line=ACC, line_w=3, shape=MSO_SHAPE.OVAL)
    text(slide, cx - d / 2, cy - d / 2, d, d, label, size=size, color=ACC, bold=True, align='c', anchor='m', lh=1.1)


def stat(slide, x, y, w, num, cap, size=96, color=ACC, cap_size=S_SMALL + 2, note=None):
    """一组数字：大号数字 + 一行说明（+ 可选灰色注）。"""
    text(slide, x, y, w, size * 1.08, num, size=size, color=color, bold=True, lh=1.0)
    text(slide, x, y + size * 1.1, w, 40, cap, size=cap_size, alpha=A_SUB, lh=1.35)
    if note:
        text(slide, x, y + size * 1.1 + 38, w, 36, note, size=S_SMALL, alpha=A_FOOT, lh=1.35)
